import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def set_cell_background(cell, fill_hex):
    """Set background color of a table cell."""
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tc_pr.append(shd)

def set_cell_margins(cell, top=120, bottom=120, left=160, right=160):
    """Set internal padding for a table cell."""
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tc_pr.append(tc_mar)

def add_callout(doc, text, title="IMPORTANT NOTICE", border_color="0D9488", bg_color="F0FDFA"):
    """Adds an executive callout box with a colored left border and light background."""
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Inches(6.8)

    cell = table.cell(0, 0)
    set_cell_background(cell, bg_color)
    set_cell_margins(cell, top=140, bottom=140, left=200, right=200)

    tc_pr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:top w:val="none"/>'
        f'<w:left w:val="single" w:sz="24" w:space="0" w:color="{border_color}"/>'
        f'<w:bottom w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'</w:tcBorders>'
    )
    tc_pr.append(borders)

    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15

    r_title = p.add_run(f"[{title}] ")
    r_title.bold = True
    r_title.font.size = Pt(9.5)
    r_title.font.name = 'Segoe UI'

    if border_color == "E11D48": # Red
        r_title.font.color.rgb = RGBColor(190, 18, 60)
    elif border_color == "0D9488": # Teal
        r_title.font.color.rgb = RGBColor(15, 118, 110)
    elif border_color == "D97706": # Amber
        r_title.font.color.rgb = RGBColor(180, 83, 9)
    elif border_color == "2563EB": # Blue
        r_title.font.color.rgb = RGBColor(37, 99, 235)
    else:
        r_title.font.color.rgb = RGBColor(30, 41, 59)

    r_body = p.add_run(text)
    r_body.font.size = Pt(9.5)
    r_body.font.name = 'Segoe UI'
    r_body.font.color.rgb = RGBColor(51, 65, 85)

    p_after = doc.add_paragraph()
    p_after.paragraph_format.space_before = Pt(0)
    p_after.paragraph_format.space_after = Pt(6)

def add_styled_table(doc, headers, rows, col_widths=None):
    """Creates a styled table with dark header and alternating row colors."""
    table = doc.add_table(rows=len(rows) + 1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    # Header row
    for i, title in enumerate(headers):
        cell = table.cell(0, i)
        set_cell_background(cell, "0F172A") # Navy slate
        set_cell_margins(cell, top=120, bottom=120, left=140, right=140)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(title)
        r.bold = True
        r.font.size = Pt(9)
        r.font.name = 'Segoe UI'
        r.font.color.rgb = RGBColor(255, 255, 255)

    # Data rows
    for r_idx, row_data in enumerate(rows):
        bg = "F8FAFC" if r_idx % 2 == 1 else "FFFFFF"
        for c_idx, val in enumerate(row_data):
            cell = table.cell(r_idx + 1, c_idx)
            set_cell_background(cell, bg)
            set_cell_margins(cell, top=100, bottom=100, left=140, right=140)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            r = p.add_run(str(val))
            r.font.size = Pt(8.5)
            r.font.name = 'Segoe UI'
            r.font.color.rgb = RGBColor(30, 41, 59)

    if col_widths:
        for row in table.rows:
            for idx, width in enumerate(col_widths):
                row.cells[idx].width = Inches(width)

    p_after = doc.add_paragraph()
    p_after.paragraph_format.space_before = Pt(0)
    p_after.paragraph_format.space_after = Pt(6)

def add_heading_1(doc, text):
    h = doc.add_paragraph()
    h.paragraph_format.space_before = Pt(16)
    h.paragraph_format.space_after = Pt(4)
    h.paragraph_format.keep_with_next = True
    r = h.add_run(text)
    r.font.name = 'Segoe UI'
    r.font.size = Pt(14)
    r.bold = True
    r.font.color.rgb = RGBColor(15, 23, 42) # Slate-900

def add_heading_2(doc, text):
    h = doc.add_paragraph()
    h.paragraph_format.space_before = Pt(12)
    h.paragraph_format.space_after = Pt(3)
    h.paragraph_format.keep_with_next = True
    r = h.add_run(text)
    r.font.name = 'Segoe UI'
    r.font.size = Pt(11.5)
    r.bold = True
    r.font.color.rgb = RGBColor(13, 148, 136) # Teal-600

def add_heading_3(doc, text):
    h = doc.add_paragraph()
    h.paragraph_format.space_before = Pt(8)
    h.paragraph_format.space_after = Pt(2)
    h.paragraph_format.keep_with_next = True
    r = h.add_run(text)
    r.font.name = 'Segoe UI'
    r.font.size = Pt(10)
    r.bold = True
    r.font.color.rgb = RGBColor(71, 85, 105) # Slate-600

def add_bullet_point(doc, title, text):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.15
    r1 = p.add_run(title + ": ")
    r1.bold = True
    r1.font.name = 'Segoe UI'
    r1.font.size = Pt(9.5)
    r1.font.color.rgb = RGBColor(30, 41, 59)
    r2 = p.add_run(text)
    r2.font.name = 'Segoe UI'
    r2.font.size = Pt(9.5)
    r2.font.color.rgb = RGBColor(51, 65, 85)

def build_quiz_user_guide(output_path):
    doc = docx.Document()

    # Page Margins (0.75 in)
    for section in doc.sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)

        # Header
        header = section.header
        hp = header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hrun = hp.add_run("Greg ICT Hub LMS • TVET MCQ Examination & Assessment Engine Operational Manual")
        hrun.font.name = 'Segoe UI'
        hrun.font.size = Pt(8.5)
        hrun.font.color.rgb = RGBColor(148, 163, 184)

        # Footer
        footer = section.footer
        fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        frun = fp.add_run("© 2026 Greg & Co. ICT Ltd • Benue State IDEAS-TVET Project • https://lms.gregicthub.com")
        frun.font.name = 'Segoe UI'
        frun.font.size = Pt(8)
        frun.font.color.rgb = RGBColor(148, 163, 184)

    # Document Styles
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Segoe UI'
    normal_style.font.size = Pt(10)
    normal_style.font.color.rgb = RGBColor(30, 41, 59)
    normal_style.paragraph_format.line_spacing = 1.15
    normal_style.paragraph_format.space_after = Pt(6)

    # Cover / Header Box
    head_table = doc.add_table(rows=1, cols=1)
    head_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    head_table.autofit = False
    head_table.columns[0].width = Inches(6.8)

    hcell = head_table.cell(0, 0)
    set_cell_background(hcell, "0F172A") # Dark Navy
    set_cell_margins(hcell, top=260, bottom=260, left=300, right=300)

    hp1 = hcell.paragraphs[0]
    hp1.alignment = WD_ALIGN_PARAGRAPH.LEFT
    hr1 = hp1.add_run("BENUE STATE IDEAS-TVET CAPACITY BUILDING INITIATIVE\nGREG & CO. ICT HUB LEARNING MANAGEMENT SYSTEM (LMS)")
    hr1.font.name = 'Segoe UI'
    hr1.font.size = Pt(9.5)
    hr1.font.bold = True
    hr1.font.color.rgb = RGBColor(45, 212, 191) # Teal

    hp2 = hcell.add_paragraph()
    hp2.paragraph_format.space_before = Pt(8)
    hr2 = hp2.add_run("TVET MCQ Examination & Competency Assessment Module\nComprehensive User & Administrator Operational Manual")
    hr2.font.name = 'Segoe UI'
    hr2.font.size = Pt(16)
    hr2.font.bold = True
    hr2.font.color.rgb = RGBColor(255, 255, 255)

    hp3 = hcell.add_paragraph()
    hp3.paragraph_format.space_before = Pt(6)
    hr3 = hp3.add_run("Official Technical Reference for Candidates, TVET Teachers, Evaluators, and Platform Administrators")
    hr3.font.name = 'Segoe UI'
    hr3.font.size = Pt(10)
    hr3.font.color.rgb = RGBColor(203, 213, 225)

    hp4 = hcell.add_paragraph()
    hp4.paragraph_format.space_before = Pt(10)
    hr4 = hp4.add_run("Version: 1.0 Production Release • Published: October 2026 • Makurdi, Benue State, Nigeria")
    hr4.font.name = 'Segoe UI'
    hr4.font.size = Pt(8.5)
    hr4.font.color.rgb = RGBColor(148, 163, 184)

    p_spacer = doc.add_paragraph()
    p_spacer.paragraph_format.space_before = Pt(8)

    # 1. Executive Summary & Capabilities
    add_heading_1(doc, "1. Executive Overview & System Architecture")
    p = doc.add_paragraph()
    p.add_run("The ")
    r = p.add_run("TVET MCQ Examination & Competency Assessment Engine")
    r.bold = True
    p.add_run(
        " is a mission-critical digital assessment platform within the Greg ICT Hub LMS. "
        "It provides rigorous, competency-oriented multiple-choice question testing for technical teachers, vocational trainees, "
        "and craft apprentices across participating Technical Colleges in Benue State and partner institutions."
    )

    add_callout(
        doc,
        "All assessments feature server-side anti-leak mechanisms: correct option keys and rationales are kept completely "
        "hidden from client browser inspection during the active exam session and are only revealed upon finalized submission.",
        title="SECURITY & INTEGRITY ASSURANCE",
        border_color="0D9488",
        bg_color="F0FDFA"
    )

    # Capabilities Table
    add_heading_2(doc, "System Role Capabilities Matrix")
    cap_headers = ["Feature / Capability", "Teacher / Candidate", "Facilitator / Evaluator", "LMS Administrator"]
    cap_rows = [
        ["View Available Exams & Instructions", "Yes (Active Only)", "Yes (All)", "Yes (All)"],
        ["Take Timed Assessment & Auto-Save", "Yes", "Yes (Practice)", "Yes (Practice)"],
        ["Immediate Result & Diagnostic Breakdown", "Yes (Own Attempt)", "Yes (All Teachers)", "Yes (All Teachers)"],
        ["Print Official Individual Evaluation Slip", "Yes (Personal)", "Yes (Any / Batch)", "Yes (Any / Batch)"],
        ["Institutional Analytics & Pass Rate KPI", "No", "Yes (Real-Time)", "Yes (Real-Time)"],
        ["Multi-Filter by State, College, Status", "No", "Yes", "Yes"],
        ["Collective Broadsheet Roster Printing", "No", "Yes", "Yes"],
        ["Streaming UTF-8 CSV Score Export", "No", "Yes", "Yes"],
        ["Create New Examinations & Categories", "No", "Yes", "Yes"],
        ["Individual Question Authoring & Editing", "No", "Yes", "Yes"],
        ["Bulk 9-Column CSV Question Ingestion", "No", "Yes (Append/Replace)", "Yes (Append/Replace)"],
        ["Purge / Cascade Delete Exam & Attempts", "No", "No", "Yes (Admin Only)"]
    ]
    add_styled_table(doc, cap_headers, cap_rows, [2.3, 1.5, 1.5, 1.5])

    # 2. Candidate & Teacher Guide
    add_heading_1(doc, "2. Candidate & Teacher Guide: Taking Assessments")
    p2 = doc.add_paragraph()
    p2.add_run(
        "Examinations are engineered to offer a seamless, responsive, and stress-free user experience. "
        "Below are the operational steps required for candidates sitting for competency assessments."
    )

    add_heading_2(doc, "Step-by-Step Assessment Workflow")
    add_bullet_point(doc, "Navigation", "Navigate to Exams & MCQs on the top menu and select Available Competency Exams (/quizzes).")
    add_bullet_point(doc, "Exam Selection", "Review available exams. Each card details the title, trade specialization, allocated time, passing mark (%), and total questions.")
    add_bullet_point(doc, "Past Attempt Audits", "If you have previously taken the test, your past scores, pass/fail status, and timestamp are displayed directly on the card.")
    add_bullet_point(doc, "Launching the Exam", "Click 'Start Assessment' (or 'Retake Assessment'). The exam environment opens with instructions and safety rules.")

    add_heading_2(doc, "Assessment Interface Features")
    add_bullet_point(doc, "High-Contrast Timer", "Displays at the top right of the screen. Color shifts to amber when 5 minutes remain and pulses red during the final 60 seconds.")
    add_bullet_point(doc, "Question Jump Drawer", "A right-hand interactive sidebar displays all question numbers. Completed questions turn solid green, while unanswered items remain amber.")
    add_bullet_point(doc, "Auto-State Preservation", "Every chosen radio button is saved instantaneously in browser memory. Even if your browser tab is refreshed accidentally, your selections are retained.")
    add_bullet_point(doc, "Submission Safeguard", "Clicking 'Submit Assessment' prompts a confirmation modal stating how many questions have been answered vs left blank.")

    add_heading_2(doc, "Official TVET Grading Scale & Evaluation Slip")
    p_grade = doc.add_paragraph()
    p_grade.add_run("Upon submission, scores are computed instantly against the official Benue State IDEAS-TVET competency scale:")

    grade_headers = ["Grading Band", "Percentage Score", "Competency Determination", "Official Recognition"]
    grade_rows = [
        ["Distinction", "80% – 100%", "Competent with Mastery", "Master Technical Instructor / Exemplary Portfolio"],
        ["Merit", "70% – 79.9%", "Competent with Commendation", "Advanced Instructional Delivery Capability"],
        ["Credit", "60% – 69.9%", "Competent", "Meets National TVET Pedagogy Standards"],
        ["Pass", "50% – 59.9%", "Bare Competent", "Satisfactory Basic Classroom Delivery"],
        ["Needs Improvement", "Below 50%", "Not Yet Competent", "Mandatory Workshop Remediation Required"]
    ]
    add_styled_table(doc, grade_headers, grade_rows, [1.5, 1.3, 1.8, 2.2])

    add_callout(
        doc,
        "Click the 'Print Official Evaluation Slip' button on your results screen to generate a signed, "
        "print-formatted evaluation sheet containing your verified identity, institution, itemized score, "
        "and supervisor signature line suitable for your physical TVET Teaching Portfolio.",
        title="PORTFOLIO PRINTING",
        border_color="2563EB",
        bg_color="EFF6FF"
    )

    # 3. Facilitator & Evaluator Guide
    add_heading_1(doc, "3. Facilitator & Evaluator Guide: Analytics & Scorecards")
    p3 = doc.add_paragraph()
    p3.add_run(
        "Institutional facilitators and evaluators hold real-time supervisory access to the "
        "Results Analytics Dashboard (/admin/quiz-analytics or /quizzes/analytics)."
    )

    add_heading_2(doc, "Institutional Evaluation Metrics")
    add_bullet_point(doc, "Total Submissions", "Monitors overall submission volume across all participating colleges.")
    add_bullet_point(doc, "Unique Teachers Evaluated", "De-duplicates attempts to indicate exactly how many distinct instructors have participated.")
    add_bullet_point(doc, "Institutional Passing Rate", "Provides macro-level percentage tracking of competency achievement.")
    add_bullet_point(doc, "Score Range Benchmarks", "Tracks highest, lowest, and mean cohort percentage scores.")

    add_heading_2(doc, "Multi-Dimensional Filtering & Cohort Analysis")
    p_filter = doc.add_paragraph()
    p_filter.add_run(
        "Evaluators can pinpoint training needs using the multi-criteria filter toolbar: "
        "filter by State (e.g. Benue, Plateau), specific Technical College / Institution, "
        "Competency Status (Passed vs Needs Improvement), or search directly by teacher name or trade area."
    )

    add_heading_2(doc, "Question Gap & Misconception Analysis")
    add_callout(
        doc,
        "The analytics engine automatically computes which 5 assessment items exhibited the highest error rates across the cohort. "
        "Facilitators must inspect this section before practical training sessions to identify pedagogical gaps "
        "(e.g., Bloom's taxonomy misapplication, workshop safety standards, or lesson timing calculations).",
        title="INSTRUCTIONAL GAP INTERVENTION",
        border_color="D97706",
        bg_color="FFFBEB"
    )

    add_heading_2(doc, "Roster Broadsheet & Batch Evaluation Slips")
    add_bullet_point(doc, "Collective Broadsheet Roster", "Generates a printable master institutional broadsheet with candidate details, scores, and institutional verification blocks.")
    add_bullet_point(doc, "Batch Individual Printing", "Select multiple teachers using the checkboxes, then click 'Print Selected Slips' to produce consecutive, page-broken official slips.")
    add_bullet_point(doc, "Streaming CSV Export", "Click 'Export CSV' to stream an Excel-compatible, BOM-encoded UTF-8 spreadsheet with complete metadata.")

    # 4. Administrator Guide: Exam & Question Bank Management
    add_heading_1(doc, "4. Platform Administrator Guide: Exam Creation & Bulk CSV Ingestion")
    p4 = doc.add_paragraph()
    p4.add_run(
        "Platform Administrators possess full administrative rights to author new examinations, modify passing thresholds, "
        "manage individual questions, and perform bulk CSV question uploads."
    )

    add_heading_2(doc, "Creating New Competency Examinations")
    add_bullet_point(doc, "Initiating Creation", "From the Available Exams screen (/quizzes), click the green '+ Create New Exam' button.")
    add_bullet_point(doc, "Exam Title", "Provide a standardized descriptive title (e.g., 'Solar PV Installation Competency Examination').")
    add_bullet_point(doc, "Lecture / Trade Category", "Select from 20+ built-in trade disciplines or click '+ Custom Category' to define a new specialization.")
    add_bullet_point(doc, "Duration & Pass Mark", "Configure time limits in minutes (e.g., 30, 45, 60) and minimum pass percentage (e.g., 50%, 65%, 75%).")
    add_bullet_point(doc, "Active / Draft State", "Toggle publish status. Draft exams remain hidden from ordinary candidates until published.")

    add_heading_2(doc, "Question Bank & Bulk CSV Upload Specification")
    p_csv = doc.add_paragraph()
    p_csv.add_run(
        "To rapidly ingest 50, 100, or more questions into any exam, use the Bulk CSV Upload tool. "
        "The uploaded CSV file must adhere strictly to the following 9-column header layout:"
    )

    csv_headers = ["Column Header", "Mandatory", "Format / Allowed Values", "Practical Example"]
    csv_rows = [
        ["category", "YES", "Category or Trade Area", "Workshop Safety & OSHA Standards"],
        ["question_text", "YES", "Full question prompt", "What is the primary function of an earth leakage circuit breaker (ELCB)?"],
        ["option_a", "YES", "Text for choice A", "Detect residual current leakage to earth and trip the supply"],
        ["option_b", "YES", "Text for choice B", "Regulate electrical line voltage during motor start-up"],
        ["option_c", "YES", "Text for choice C", "Step down secondary transformer winding current"],
        ["option_d", "YES", "Text for choice D", "Protect purely against thermal overloads"],
        ["correct_option", "YES", "A, B, C, or D (Case-insensitive)", "A"],
        ["explanation", "NO", "Pedagogical explanation for review", "ELCBs detect residual current imbalances to prevent electric shock."],
        ["points", "NO", "Numeric value (Default: 1.0)", "1.0"]
    ]
    add_styled_table(doc, csv_headers, csv_rows, [1.3, 1.0, 2.0, 2.5])

    add_heading_2(doc, "Ingestion Strategies: Append vs Replace")
    add_bullet_point(doc, "Append Mode", "Maintains any existing questions in the exam and appends new items from the CSV to the end of the question bank.")
    add_bullet_point(doc, "Replace Mode", "Safely purges existing questions for that specific exam and replaces them with the contents of the CSV.")

    add_callout(
        doc,
        "The system enforces a 60MB file upload ceiling. Always ensure your CSV file is saved with UTF-8 encoding "
        "and standard comma delimiters. A pre-formatted sample CSV template is downloadable directly from the Question Bank screen.",
        title="CSV ENCODING & UPLOAD SPECS",
        border_color="0D9488",
        bg_color="F0FDFA"
    )

    # 5. Standard Operating Procedures & Maintenance
    add_heading_1(doc, "5. Standard Operating Procedures & Server Maintenance")
    p5 = doc.add_paragraph()
    p5.add_run(
        "To ensure uninterrupted reliability and prevent 404 route errors on production shared hosting environments "
        "(such as Hostinger), adhere to the following maintenance procedures."
    )

    add_heading_2(doc, "Production Deployment Command Checklist")
    p_cmd = doc.add_paragraph()
    p_cmd.add_run(
        "Execute the following command sequence inside the application directory whenever new updates or route definitions are deployed:\n"
        "  1. git pull origin main\n"
        "  2. php artisan optimize:clear\n"
        "  3. php artisan route:cache\n"
        "  4. php artisan view:cache\n"
        "  5. php artisan config:cache"
    )

    add_heading_2(doc, "Technical Support & Coordination")
    add_bullet_point(doc, "Operating Agency", "Greg & Co. ICT Ltd — Lead TVET Digital Innovations Partner")
    add_bullet_point(doc, "Target Project", "Benue State IDEAS-TVET Technical Teacher Capacity Building Initiative")
    add_bullet_point(doc, "Host Location", "Makurdi, Benue State, Nigeria")
    add_bullet_point(doc, "Production Portal", "https://lms.gregicthub.com")

    # Save Document
    doc.save(output_path)
    print(f"Document successfully created at: {output_path}")

if __name__ == '__main__':
    downloads_dir = os.path.join(os.path.dirname(__file__), 'public', 'downloads')
    os.makedirs(downloads_dir, exist_ok=True)
    target_path = os.path.join(downloads_dir, 'TVET_MCQ_Examination_Module_Comprehensive_User_Guide.docx')
    build_quiz_user_guide(target_path)
